"""Worksheet picker combo behaviour (BOM/PnP tabs)."""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook
from PySide6 import QtWidgets

from ui.sheet_picker import SHEET_AUTO, SheetPickerMixin, is_spreadsheet_path
from ui_i18n import UiI18n

pytest.importorskip("PySide6")


@pytest.fixture(scope="module")
def qt_app():
    app = QtWidgets.QApplication.instance() or QtWidgets.QApplication(sys.argv)
    yield app


@pytest.fixture
def workbook(tmp_path: Path) -> str:
    """Hidden ECN sheet (5 rows) plus visible SKU3 sheet (200 rows)."""
    wb = Workbook()
    wb.active.title = "ECN"
    for i in range(5):
        wb.active.append([i, f"R{i}"])
    sku = wb.create_sheet("SKU3")
    for i in range(200):
        sku.append([1, f"C{i}", "MLCC", f"CL{i}"])
    wb.active.sheet_state = "hidden"
    path = tmp_path / "bom.xlsx"
    wb.save(path)
    return str(path)


class _Host(SheetPickerMixin):
    """Minimal stand-in exposing just what SheetPickerMixin touches."""

    def __init__(self, app):
        widget = QtWidgets.QWidget()
        self.widget = widget  # keep alive; owns the combo layout
        self._bom_source_path = ""
        self._pnp_source_path = ""
        self._sheet_paths: dict[str, str] = {}
        self._sheets_updating = False
        self.logs: list[tuple[str, str]] = []
        self._reload_count = 0

        # Real English catalog so format placeholders match production.
        self.ui_tr = UiI18n("en").tr
        self._log = lambda msg, level="info": self.logs.append((level, msg))
        layout = QtWidgets.QVBoxLayout(widget)
        layout.addLayout(self._create_sheet_picker("bom", tr_key="bom.sheet"))

        def _reload_bom():
            self._reload_count += 1

        self._reload_bom = _reload_bom
        self._reload_pnp = _reload_bom
        self._schedule_save_bom_tab_settings = lambda: None
        self._schedule_save_pnp_tab_settings = lambda: None


@pytest.fixture
def host(qt_app, workbook) -> _Host:
    h = _Host(qt_app)
    h._bom_source_path = workbook
    h._populate_sheet_combo("bom", workbook)
    return h


def test_combo_lists_all_sheets_with_hidden_marker(host: _Host) -> None:
    combo = host.bom_sheet
    assert combo.count() == 3
    assert combo.itemText(0) == SHEET_AUTO
    labels = [combo.itemText(i) for i in range(combo.count())]
    assert any(l.endswith("(hidden)") for l in labels)
    assert host._combo_index_for_sheet(combo, "SKU3") >= 0


def test_auto_selects_first_visible_sheet(host: _Host) -> None:
    assert host.bom_sheet.currentText() == "SKU3"
    assert host._selected_sheet("bom") == "SKU3"


def test_csv_disables_combo(qt_app, tmp_path: Path) -> None:
    h = _Host(qt_app)
    csv_path = tmp_path / "bom.csv"
    csv_path.write_text("Ref,Value\nR1,10k\n", encoding="utf-8")
    h._populate_sheet_combo("bom", str(csv_path))
    assert not h.bom_sheet.isEnabled()
    assert h._selected_sheet("bom") is None
    assert h.bom_sheet_rows.text() == ""


def test_missing_path_disables_combo(qt_app) -> None:
    h = _Host(qt_app)
    h._populate_sheet_combo("bom", "")
    assert not h.bom_sheet.isEnabled()


def test_hidden_sheet_can_be_selected_explicitly(host: _Host) -> None:
    idx = host._combo_index_for_sheet(host.bom_sheet, "ECN")
    host.bom_sheet.setCurrentIndex(idx)
    assert host._selected_sheet("bom") == "ECN"


def test_stale_saved_sheet_warns_and_falls_back(host: _Host) -> None:
    """A remembered sheet that vanished warns instead of failing the load."""
    host.bom_sheet.blockSignals(True)
    host.bom_sheet.setCurrentIndex(0)
    host.bom_sheet.blockSignals(False)
    host._populate_sheet_combo("bom", host._bom_source_path, saved_sheet="GONE")
    assert host._selected_sheet("bom") == "SKU3"
    assert any(level == "warning" for level, _ in host.logs)


def test_populate_does_not_trigger_reload(host: _Host) -> None:
    assert host._reload_count == 0


def test_current_sheet_survives_repopulation(host: _Host) -> None:
    """Reload repopulates the combo; the live choice must not snap back."""
    idx = host._combo_index_for_sheet(host.bom_sheet, "ECN")
    host.bom_sheet.setCurrentIndex(idx)
    host._reload_count = 0
    host._populate_sheet_combo("bom", host._bom_source_path)
    assert host._selected_sheet("bom") == "ECN"


def test_selection_does_not_leak_between_files(qt_app, workbook, tmp_path) -> None:
    """Another workbook starts from its own default, not the previous choice."""
    h = _Host(qt_app)
    h._bom_source_path = workbook
    h._populate_sheet_combo("bom", workbook)
    idx = h._combo_index_for_sheet(h.bom_sheet, "ECN")
    h.bom_sheet.setCurrentIndex(idx)

    other = tmp_path / "other.xlsx"
    wb = Workbook()
    wb.active.title = "ONLY"
    wb.active.append(["a"])
    wb.active.append(["b"])
    wb.save(other)
    h._bom_source_path = str(other)
    h._populate_sheet_combo("bom", str(other))
    assert h._selected_sheet("bom") == "ONLY"


def test_clear_resets_picker(qt_app, workbook) -> None:
    h = _Host(qt_app)
    h._populate_sheet_combo("bom", workbook)
    h._clear_sheet_picker("bom")
    assert h.bom_sheet.currentText() == SHEET_AUTO
    assert not h.bom_sheet.isEnabled()
    assert h._selected_sheet("bom") is None


def test_hint_shows_sheet_and_rows(qt_app, workbook) -> None:
    h = _Host(qt_app)
    h._populate_sheet_combo("bom", workbook)
    h._set_sheet_hint("bom", "SKU3", 200)
    assert "SKU3" in h.bom_sheet_rows.text()
    assert "200" in h.bom_sheet_rows.text()


def test_hint_without_rows_shows_sheet_name(qt_app, workbook) -> None:
    h = _Host(qt_app)
    h._populate_sheet_combo("bom", workbook)
    h._set_sheet_hint("bom", "SKU3", None)
    assert h.bom_sheet_rows.text() == "SKU3"


def test_hint_empty_for_blank_sheet(qt_app, workbook) -> None:
    h = _Host(qt_app)
    h._populate_sheet_combo("bom", workbook)
    h._set_sheet_hint("bom", "", None)
    assert h.bom_sheet_rows.text() == ""


def test_user_selection_requests_reload(host: _Host) -> None:
    idx = host._combo_index_for_sheet(host.bom_sheet, "ECN")
    host.bom_sheet.setCurrentIndex(idx)
    assert host._reload_count == 1


def test_combo_index_unknown_sheet_is_minus_one(host: _Host) -> None:
    assert host._combo_index_for_sheet(host.bom_sheet, "NOPE") == -1


def test_combo_index_none_is_auto(host: _Host) -> None:
    assert host._combo_index_for_sheet(host.bom_sheet, None) == 0


def test_is_spreadsheet_path() -> None:
    assert is_spreadsheet_path("a.xlsx")
    assert is_spreadsheet_path("a.XLS")
    assert is_spreadsheet_path("a.xlsm")
    assert is_spreadsheet_path("a.ods")
    assert not is_spreadsheet_path("a.csv")
    assert not is_spreadsheet_path("a.txt")
