"""Worksheet selection: first-visible default, explicit sheet, missing sheet.

Regression cover for the TechOne 27 (CO1271) BOM, whose first sheet
(ECN) was hidden while the visible SKU3 sheet held the real 266 rows.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

import smt_processor

_ASSETS = Path(__file__).resolve().parent / "assets"


def _multi_sheet_workbook(
    path: Path,
    *,
    hidden_first: bool,
    hidden_last: bool = False,
) -> str:
    """BOM workbook: hidden ECN log (39 rows), visible SKU3 (266 rows), notes."""
    wb = Workbook()
    wb.active.title = "ECN"
    for i in range(39):
        wb.active.append([i, f"R{i}"])

    sku = wb.create_sheet("SKU3")
    sku.append([1, "PCBA", "MDXB101003"])
    for i in range(265):
        sku.append([1, f"C{i}", f"MLCC_{i}", f"CL{i}"])

    notes = wb.create_sheet("NOTES")
    notes.append(["production note"])

    wb.active.sheet_state = "hidden" if hidden_first else "visible"
    wb["NOTES"].sheet_state = "hidden" if hidden_last else "visible"
    wb.save(path)
    return str(path)


def test_list_sheets_reports_hidden(tmp_path: Path) -> None:
    path = _multi_sheet_workbook(tmp_path / "bom.xlsx", hidden_first=True)
    sheets = smt_processor.list_sheets(path)
    assert [s.name for s in sheets] == ["ECN", "SKU3", "NOTES"]
    assert [s.visible for s in sheets] == [False, True, True]
    assert sheets[0].label.endswith("(hidden)")


def test_default_sheet_skips_hidden_first_sheet(tmp_path: Path) -> None:
    """The reported bug: hidden ECN (39) was loaded instead of visible SKU3."""
    path = _multi_sheet_workbook(tmp_path / "bom.xlsx", hidden_first=True)
    assert smt_processor.default_sheet_name(smt_processor.list_sheets(path)) == "SKU3"

    df = smt_processor.read_file(path, column_headers_from_file=False)
    assert len(df) == 266


def test_default_sheet_is_first_when_none_hidden(tmp_path: Path) -> None:
    path = _multi_sheet_workbook(tmp_path / "bom.xlsx", hidden_first=False)
    assert smt_processor.default_sheet_name(smt_processor.list_sheets(path)) == "ECN"
    assert len(smt_processor.read_file(path, column_headers_from_file=False)) == 39


def test_default_sheet_falls_back_when_all_hidden() -> None:
    """No visible sheet: still pick something instead of failing.

    Exercised on the SheetInfo list directly - openpyxl refuses to save a
    workbook with every sheet hidden, so no real file can represent this.
    """
    sheets = [
        smt_processor.SheetInfo("A", False),
        smt_processor.SheetInfo("B", False),
    ]
    assert smt_processor.default_sheet_name(sheets) == "A"


def test_explicit_sheet_name_reads_that_sheet(tmp_path: Path) -> None:
    path = _multi_sheet_workbook(tmp_path / "bom.xlsx", hidden_first=True)
    df = smt_processor.read_file(path, sheet_name="ECN", column_headers_from_file=False)
    assert len(df) == 39


def test_explicit_sheet_overrides_default(tmp_path: Path) -> None:
    path = _multi_sheet_workbook(tmp_path / "bom.xlsx", hidden_first=True)
    df = smt_processor.read_file(
        path, sheet_name="NOTES", column_headers_from_file=False
    )
    assert len(df) == 1


def test_missing_sheet_raises_sheet_not_found(tmp_path: Path) -> None:
    """A bad sheet must not be masked by the CSV fallback path."""
    path = _multi_sheet_workbook(tmp_path / "bom.xlsx", hidden_first=True)
    with pytest.raises(smt_processor.SMTSheetNotFoundError):
        smt_processor.read_file(path, sheet_name="NO_SUCH_SHEET")


def test_sheet_not_found_is_processor_error(tmp_path: Path) -> None:
    """Callers catching SMTProcessorError keep working."""
    path = _multi_sheet_workbook(tmp_path / "bom.xlsx", hidden_first=True)
    with pytest.raises(smt_processor.SMTProcessorError):
        smt_processor.read_file(path, sheet_name="NO_SUCH_SHEET")


def test_list_sheets_empty_for_csv(tmp_path: Path) -> None:
    path = tmp_path / "bom.csv"
    path.write_text("Designator,Comment\nR1,10k\n", encoding="utf-8")
    assert smt_processor.list_sheets(str(path)) == []


def test_list_sheets_empty_for_missing_file() -> None:
    assert smt_processor.list_sheets("no_such_file.xlsx") == []


def test_default_sheet_name_empty_list() -> None:
    assert smt_processor.default_sheet_name([]) is None


@pytest.mark.parametrize("suffix", [".xls", ".ods"])
def test_single_sheet_asset_first_sheet_is_readable(suffix: str) -> None:
    """Legacy .xls/.ods paths still resolve a sheet."""
    path = _ASSETS / f"bom{suffix}"
    if not path.is_file():
        pytest.skip(f"bom{suffix} missing")
    sheets = smt_processor.list_sheets(str(path))
    assert sheets
    default = smt_processor.default_sheet_name(sheets)
    assert default == sheets[0].name
    assert len(smt_processor.read_file(str(path))) > 0
